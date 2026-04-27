"""
Slide Manager — loads PPTX / PDF, navigates slides, exports PNG previews.
"""
import logging
import os
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class SlideData:
    index: int          # 0-based
    total: int
    title: str
    content: str
    speaker_notes: str

    @property
    def number(self) -> int:
        return self.index + 1

    def summary(self) -> str:
        """Compact text representation for the AI prompt."""
        parts = [f"Slide {self.number}/{self.total}: {self.title}"]
        if self.content:
            parts.append(self.content)
        if self.speaker_notes:
            parts.append(f"[Notes: {self.speaker_notes}]")
        return "\n".join(parts)


class SlideManager:
    def __init__(self, image_dir: str = "slide_images"):
        self._slides: list[SlideData] = []
        self._index: int = 0
        self._source: Optional[str] = None
        self._image_dir = image_dir

    # ------------------------------------------------------------------ #
    # Loading
    # ------------------------------------------------------------------ #

    def load(self, path: str) -> int:
        self._source = path
        self._slides = []
        self._index = 0

        ext = Path(path).suffix.lower()
        if ext in (".pptx", ".ppt"):
            self._load_pptx(path)
        elif ext == ".pdf":
            self._load_pdf(path)
        else:
            raise ValueError(f"Unsupported format: {ext}")

        total = len(self._slides)
        for s in self._slides:
            s.total = total
        logger.info("Loaded %d slides from %s", total, path)
        return total

    def _load_pptx(self, path: str):
        from pptx import Presentation

        prs = Presentation(path)
        for i, slide in enumerate(prs.slides):
            title, content = "", []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                name = shape.name.lower()
                is_title = "title" in name or (
                    hasattr(shape, "placeholder_format")
                    and shape.placeholder_format is not None
                    and shape.placeholder_format.idx == 0
                )
                if is_title:
                    title = text
                else:
                    content.append(text)
            notes = ""
            if slide.has_notes_slide:
                try:
                    nf = slide.notes_slide.notes_text_frame
                    notes = nf.text.strip() if nf else ""
                except Exception:
                    pass
            self._slides.append(SlideData(i, 0, title or f"Slide {i+1}", "\n".join(content), notes))

    def _load_pdf(self, path: str):
        import fitz

        doc = fitz.open(path)
        for i in range(len(doc)):
            text = doc[i].get_text().strip()
            lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
            title = lines[0] if lines else f"Slide {i+1}"
            content = "\n".join(lines[1:]) if len(lines) > 1 else ""
            self._slides.append(SlideData(i, 0, title, content, ""))
        doc.close()

    # ------------------------------------------------------------------ #
    # Navigation
    # ------------------------------------------------------------------ #

    @property
    def total(self) -> int:
        return len(self._slides)

    @property
    def is_loaded(self) -> bool:
        return bool(self._slides)

    @property
    def at_end(self) -> bool:
        return self._index >= len(self._slides) - 1

    @property
    def at_start(self) -> bool:
        return self._index == 0

    def current(self) -> Optional[SlideData]:
        return self._slides[self._index] if self._slides else None

    def next(self) -> Optional[SlideData]:
        if not self.at_end:
            self._index += 1
        return self.current()

    def previous(self) -> Optional[SlideData]:
        if not self.at_start:
            self._index -= 1
        return self.current()

    def goto(self, index: int) -> Optional[SlideData]:
        self._index = max(0, min(index, len(self._slides) - 1))
        return self.current()

    # ------------------------------------------------------------------ #
    # PNG export
    # ------------------------------------------------------------------ #

    def export_image(self, index: int) -> Optional[str]:
        """Render slide to PNG. Returns path or None on failure."""
        if not self._source:
            return None
        os.makedirs(self._image_dir, exist_ok=True)
        out = os.path.join(self._image_dir, f"slide_{index:03d}.png")
        if os.path.exists(out):
            return out

        ext = Path(self._source).suffix.lower()
        if ext == ".pdf":
            return self._render_pdf(self._source, index, out)
        elif ext in (".pptx", ".ppt"):
            pdf = self._pptx_to_pdf(self._source)
            return self._render_pdf(pdf, index, out) if pdf else None
        return None

    def export_all_images(self):
        for i in range(self.total):
            self.export_image(i)

    def _render_pdf(self, pdf: str, page: int, out: str) -> Optional[str]:
        try:
            import fitz
            doc = fitz.open(pdf)
            if page >= len(doc):
                return None
            pix = doc[page].get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
            pix.save(out)
            doc.close()
            return out
        except Exception as exc:
            logger.warning("PDF render failed (page %d): %s", page, exc)
            return None

    def _pptx_to_pdf(self, pptx: str) -> Optional[str]:
        pdf = pptx.replace(".pptx", ".pdf").replace(".ppt", ".pdf")
        if os.path.exists(pdf):
            return pdf
        try:
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "pdf",
                 "--outdir", os.path.dirname(pptx), pptx],
                check=True, capture_output=True, timeout=60,
            )
            return pdf if os.path.exists(pdf) else None
        except Exception as exc:
            logger.warning("LibreOffice conversion failed: %s", exc)
            return None
